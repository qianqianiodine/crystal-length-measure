"""导出标签 API（`POST /api/images/export-labels`）。

⚠️ 必须 monkeypatch config 里的目录，否则会往用户真实的 data/ 和 导出/ 里写东西。
⚠️ `os.startfile` 要挡住 —— 不加的话每跑一条测试，桌面上真的弹一个资源管理器窗口。
"""
import os
from pathlib import Path

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import config, main


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(config, "IMAGES_DIR", tmp_path / "data" / "images" / "originals")
    monkeypatch.setattr(config, "CACHE_DIR", tmp_path / "data" / "images" / "cache")
    monkeypatch.setattr(config, "THUMBS_DIR", tmp_path / "data" / "thumbs")
    monkeypatch.setattr(config, "STAGING_DIR", tmp_path / "data" / "staging")
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "data" / "app.db")
    monkeypatch.setattr(config, "EXPORT_DIR", tmp_path / "导出")
    monkeypatch.setattr(os, "startfile", lambda p: None, raising=False)
    config.ensure_dirs()
    with TestClient(main.app) as c:
        yield c


def _png(w=40, h=30) -> bytes:
    ok, buf = cv2.imencode(".png", np.full((h, w, 3), 200, np.uint8))
    assert ok
    return buf.tobytes()


def _add(client, name) -> int:
    r = client.post("/api/images/upload",
                    files=[("files", (f"{name}.png", _png(), "image/png"))])
    assert r.status_code == 200, r.text
    return r.json()["ok"][0]["id"]


def _export(client, body):
    r = client.post("/api/images/export-labels", json=body)
    assert r.status_code == 200, r.text
    j = r.json()
    return j, Path(j["file"])


def test_导出_xlsx_两列并带上完整文件夹路径(client):
    a = client.post("/api/folders", json={"name": "2026", "parent_id": None}).json()
    b = client.post("/api/folders", json={"name": "20261020_TREX1_16",
                                          "parent_id": a["id"]}).json()
    i1 = _add(client, "50-JBScreen-H-9-1")
    i2 = _add(client, "50-JBScreen-G-3-1")
    client.post("/api/images/batch-folders",
                json={"image_ids": [i1, i2], "folder_ids": [b["id"]],
                      "mode": "replace"})

    j, f = _export(client, {"image_ids": [i1, i2], "format": "xlsx",
                            "with_folders": True})
    assert j["count"] == 2 and j["failed"] == []
    assert f.is_file() and f.suffix == ".xlsx"

    ws = load_workbook(f).active
    assert [c.value for c in ws[1]] == ["标签", "文件夹"]
    assert [ws["A2"].value, ws["B2"].value] == [
        "50-JBScreen-H-9-1", "2026/20261020_TREX1_16"]
    assert [ws["A3"].value, ws["B3"].value] == [
        "50-JBScreen-G-3-1", "2026/20261020_TREX1_16"]


def test_导出_xlsx_不勾文件夹就一列(client):
    i1 = _add(client, "甲")
    j, f = _export(client, {"image_ids": [i1], "format": "xlsx",
                            "with_folders": False})
    ws = load_workbook(f).active
    assert [c.value for c in ws[1]] == ["标签"]
    assert ws.max_column == 1


def test_导出_txt_带文件夹时文件夹名只出现一次(client):
    fid = client.post("/api/folders", json={"name": "批次A"}).json()["id"]
    i1 = _add(client, "甲")
    i2 = _add(client, "乙")
    client.post("/api/images/batch-folders",
                json={"image_ids": [i1, i2], "folder_ids": [fid], "mode": "replace"})

    j, f = _export(client, {"image_ids": [i1, i2], "format": "txt",
                            "with_folders": True})
    text = f.read_text(encoding="utf-8-sig")
    assert text == "批次A\n甲\n乙\n"
    assert text.count("批次A") == 1


def test_导出_txt_不勾文件夹就是名单(client):
    i1 = _add(client, "甲")
    i2 = _add(client, "乙")
    j, f = _export(client, {"image_ids": [i1, i2], "format": "txt",
                            "with_folders": False})
    assert f.read_text(encoding="utf-8-sig") == "甲\n乙\n"


def test_没归档的照片在清单里写未分类(client):
    i1 = _add(client, "甲")
    j, f = _export(client, {"image_ids": [i1], "format": "txt",
                            "with_folders": True})
    assert f.read_text(encoding="utf-8-sig") == "未分类\n甲\n"


def test_不存在的_id跳过其余照常(client):
    i1 = _add(client, "甲")
    j, f = _export(client, {"image_ids": [i1, 99999], "format": "txt",
                            "with_folders": False})
    assert j["count"] == 1 and j["failed"] == [99999]


def test_空列表报400(client):
    r = client.post("/api/images/export-labels",
                    json={"image_ids": [], "format": "txt"})
    assert r.status_code == 400, r.text


def test_格式不对报400(client):
    i1 = _add(client, "甲")
    r = client.post("/api/images/export-labels",
                    json={"image_ids": [i1], "format": "csv"})
    assert r.status_code == 400, r.text


def test_全都不存在也是400(client):
    r = client.post("/api/images/export-labels",
                    json={"image_ids": [99998, 99999], "format": "xlsx"})
    assert r.status_code == 400, r.text


def test_导出标签不改照片名字(client):
    i1 = _add(client, "50-JBScreen-H-9-1")
    before = client.get(f"/api/image/{i1}").json()
    _export(client, {"image_ids": [i1], "format": "xlsx", "with_folders": True})
    assert client.get(f"/api/image/{i1}").json() == before


def test_同一秒导出两次不会覆盖前一个文件(client):
    i1 = _add(client, "甲")
    j1, f1 = _export(client, {"image_ids": [i1], "format": "txt"})
    j2, f2 = _export(client, {"image_ids": [i1], "format": "txt"})
    assert f1 != f2 and f1.is_file() and f2.is_file()
