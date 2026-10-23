"""导出标签（照片名字清单）的纯函数测试：文件夹路径、txt 分组、最小 xlsx。"""
from io import BytesIO

from openpyxl import load_workbook

from app import label_export


def _img(i, name, orig="照片.png"):
    return {"id": i, "name": name, "original_filename": orig}


def _folder(i, name, parent=None):
    return {"id": i, "name": name, "parent_id": parent}


def _sheet(data: bytes):
    """拿 openpyxl 把自家拼的 xlsx 读回来 —— 读得开、格子对，才算真的能用。"""
    return load_workbook(BytesIO(data)).active


# ---------------- build_rows ----------------

def test_文件夹路径是完整路径():
    folders = [_folder(1, "2026"), _folder(2, "20261020_TREX1_16", 1)]
    rows = label_export.build_rows(
        [_img(7, "50-JBScreen-H-9-1")], folders, {7: [2]})
    assert rows == [{"label": "50-JBScreen-H-9-1",
                     "folder": "2026/20261020_TREX1_16"}]


def test_没归档的写未分类():
    rows = label_export.build_rows([_img(7, "甲")], [], {7: []})
    assert rows[0]["folder"] == label_export.UNCATEGORIZED


def test_一张图在多个文件夹里用顿号连起来():
    folders = [_folder(1, "A"), _folder(2, "B")]
    rows = label_export.build_rows([_img(7, "甲")], folders, {7: [1, 2]})
    assert rows[0]["folder"] == "A、B"


def test_文件夹数据被改坏成环也不会转死():
    folders = [_folder(1, "A", 2), _folder(2, "B", 1)]
    rows = label_export.build_rows([_img(7, "甲")], folders, {7: [1]})
    assert rows[0]["folder"]            # 有值就行，重点是没死循环


def test_名字空了退回原文件名():
    rows = label_export.build_rows([_img(7, "  ", "IMG_0007.jpg")], [], {7: []})
    assert rows[0]["label"] == "IMG_0007"


# ---------------- txt ----------------

def test_txt_不带文件夹是一行一个():
    rows = [{"label": "甲", "folder": "A"}, {"label": "乙", "folder": "B"}]
    assert label_export.to_txt(rows, False) == "甲\n乙\n"


def test_txt_带文件夹时文件夹名单独一行且只出现一次():
    rows = [{"label": "甲", "folder": "2026/批次"},
            {"label": "乙", "folder": "2026/批次"},
            {"label": "丙", "folder": "未分类"}]
    got = label_export.to_txt(rows, True)
    assert got == "2026/批次\n甲\n乙\n\n未分类\n丙\n"
    assert got.count("2026/批次") == 1


def test_txt_空清单是空字符串():
    assert label_export.to_txt([], True) == ""


# ---------------- xlsx ----------------

def test_xlsx_只有标签一列():
    ws = _sheet(label_export.to_xlsx([{"label": "甲", "folder": "A"}], False))
    assert [c.value for c in ws[1]] == ["标签"]
    assert ws["A2"].value == "甲"
    assert ws.max_column == 1


def test_xlsx_带文件夹是两列():
    rows = [{"label": "50-JBScreen-H-9-1", "folder": "2026/批次"},
            {"label": "50-JBScreen-G-3-1", "folder": "未分类"}]
    ws = _sheet(label_export.to_xlsx(rows, True))
    assert [c.value for c in ws[1]] == ["标签", "文件夹"]
    assert ws["A2"].value == "50-JBScreen-H-9-1"
    assert ws["B2"].value == "2026/批次"
    assert ws["B3"].value == "未分类"
    assert ws.max_column == 2


def test_xlsx_表头是加粗的():
    ws = _sheet(label_export.to_xlsx([], True))
    assert ws["A1"].font.bold is True


def test_xlsx_特殊字符不会把文件弄坏():
    ws = _sheet(label_export.to_xlsx(
        [{"label": "A&B <C>", "folder": 'D"E'}], True))
    assert ws["A2"].value == "A&B <C>"
    assert ws["B2"].value == 'D"E'
